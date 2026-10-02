#!/usr/bin/env python3
"""Build editable template-based papers and a companion deck from private JSON.

Usage: python build_documents.py --content private/content.json
Requires python-docx, python-pptx, lxml, latex2mathml, matplotlib, Pillow.
Official templates are supplied locally; copyrighted templates are not bundled.
"""
import argparse, json, re, io, hashlib
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.shared import Pt as DP, Cm
from docx.oxml import OxmlElement as WX
from docx.oxml.ns import qn as wqn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement as PX
from office_math import put_word, put_powerpoint

HERE=Path(__file__).resolve().parent
FONT='맑은 고딕'
BODY_FONT='바탕'

def word_font(run, size=9, bold=False, face=BODY_FONT):
    run.font.name=face;run.font.size=DP(size);run.font.bold=bold
    rf=run._r.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii','hAnsi','eastAsia']:rf.set(wqn('w:'+k),face)

def inline_word(p, text, size=9, bold=False):
    for bit in re.split(r'(\\\(.*?\\\)|\[\d+(?:,\s*\d+)*\])',text):
        if bit.startswith('\\('):put_word(p,bit[2:-2],size)
        elif re.fullmatch(r'\[\d+(?:,\s*\d+)*\]',bit):
            r=p.add_run(bit);word_font(r,size-1,bold);r.font.superscript=True
        elif bit:word_font(p.add_run(bit),size,bold)

def cols(section, count, gap):
    el=section._sectPr.find(wqn('w:cols'))
    if el is None:el=WX('w:cols');section._sectPr.append(el)
    el.set(wqn('w:num'),str(count));el.set(wqn('w:space'),str(gap))

def paper(content, template, output, blind):
    d=Document(template)
    # Retain the template styles, margins and document settings. Its sample
    # text, sample figures, title box and header/footer text are replaced.
    sect=deepcopy(d.sections[-1]._sectPr)
    for el in list(d._element.body):d._element.body.remove(el)
    d._element.body.append(sect)
    for section in d.sections:
        for part in [section.header,section.footer,section.first_page_header,section.first_page_footer]:
            for p in part.paragraphs:p.clear()
    d.core_properties.author='' if blind else content.get('author_line','')
    d.core_properties.last_modified_by=''
    d.core_properties.title=content['title_ko'];d.core_properties.subject=''
    section=d.sections[0];gap=510 if blind else 567;cols(section,1,gap)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    word_font(p.add_run(content['title_ko']),20 if blind else 18,True)
    p.paragraph_format.space_after=DP(5)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    word_font(p.add_run(content['title_en']),15,False,'Times New Roman')
    if not blind:
        for text in [content['author_line'],content['affiliation_line']]:
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            word_font(p.add_run(text),10)
    abstracts=[('요 약','abstract_ko'),('Abstract','abstract_en')] if blind else [('Abstract','abstract_en')]
    for title,key in abstracts:
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        word_font(p.add_run(title),11,True)
        p.paragraph_format.space_before=DP(3);p.paragraph_format.space_after=DP(2)
        p=d.add_paragraph();inline_word(p,content[key],9)
        p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=DP(3)
    p=d.add_paragraph();word_font(p.add_run('Keywords: '+', '.join(content['keywords'])),9,False,'Times New Roman')
    p.paragraph_format.space_after=DP(5)
    section=d.add_section(WD_SECTION_START.CONTINUOUS);cols(section,2,gap)
    last_page=1
    for block in content['paper_blocks']:
        if block.get('page',last_page)>last_page:
            # Next-page sections guarantee four actual pages instead of merely
            # inserting a column break inside a multicolumn layout.
            section=d.add_section(WD_SECTION_START.NEW_PAGE);cols(section,2,gap)
            last_page=block['page']
        kind=block['type']
        if kind=='column_break':
            d.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
        elif kind=='heading':
            p=d.add_paragraph();inline_word(p,block['text'],11,True)
            p.paragraph_format.space_before=DP(7);p.paragraph_format.space_after=DP(4)
            p.paragraph_format.keep_with_next=True
        elif kind=='subheading':
            p=d.add_paragraph();inline_word(p,block['text'],10,True)
            p.paragraph_format.space_before=DP(4);p.paragraph_format.space_after=DP(3)
            p.paragraph_format.keep_with_next=True
        elif kind=='paragraph':
            p=d.add_paragraph();inline_word(p,block['text'],9)
            p.paragraph_format.line_spacing=1.25;p.paragraph_format.space_after=DP(4)
            p.paragraph_format.first_line_indent=Cm(.18)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind=='equation':
            p=d.add_paragraph();put_word(p,block['latex'],block.get('size',9))
            word_font(p.add_run('   ('+str(block['number'])+')'),9)
            p.paragraph_format.space_before=DP(3);p.paragraph_format.space_after=DP(4)
            p.paragraph_format.keep_together=True
        elif kind=='figure':
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE/block['path']),width=Cm(block.get('width_cm',7.9)))
            p.paragraph_format.space_after=DP(2);p.paragraph_format.keep_with_next=True
            p=d.add_paragraph();inline_word(p,block['caption'],8)
            p.paragraph_format.space_after=DP(5)
        elif kind=='table':
            p=d.add_paragraph();inline_word(p,block['caption'],8)
            p.paragraph_format.keep_with_next=True
            rows=block['rows'];t=d.add_table(rows=len(rows),cols=len(rows[0]));t.autofit=False
            available_cm=7.98 if blind else 7.95
            widths=block.get('widths_cm',[available_cm/len(rows[0])]*len(rows[0]))
            for col,width in zip(t.columns,widths):col.width=Cm(width)
            t._tbl.tblPr.find(wqn('w:tblW')).set(wqn('w:w'),str(int(available_cm/2.54*1440)))
            t._tbl.tblPr.find(wqn('w:tblW')).set(wqn('w:type'),'dxa')
            for row_no,row in enumerate(rows):
                for c_no,text in enumerate(row):
                    c=t.cell(row_no,c_no);c.width=Cm(widths[c_no])
                    margins=WX('w:tcMar')
                    for side in ['left','right']:
                        side_el=WX('w:'+side);side_el.set(wqn('w:w'),'20');side_el.set(wqn('w:type'),'dxa');margins.append(side_el)
                    c._tc.get_or_add_tcPr().append(margins)
                    p=c.paragraphs[0];inline_word(p,str(text),block.get('size',7.6),row_no==0)
                    p.paragraph_format.space_after=DP(2);p.paragraph_format.space_before=DP(2)
                    p.paragraph_format.line_spacing=1.05
                trpr=t.rows[row_no]._tr.get_or_add_trPr();trpr.append(WX('w:cantSplit'))
            borders=WX('w:tblBorders')
            for side in ['top','bottom','insideH']:
                b=WX('w:'+side);b.set(wqn('w:val'),'single');b.set(wqn('w:sz'),'4');borders.append(b)
            t._tbl.tblPr.append(borders)
        elif kind=='references':
            for text in block['items']:
                p=d.add_paragraph();word_font(p.add_run(text),8,False,'Times New Roman')
                p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=DP(4)
        else:raise ValueError(kind)
    # No field, comments, author identifiers or old template images survive
    # in the body. Clear standard personal metadata for blind review.
    if blind:
        d.core_properties.keywords='';d.core_properties.comments=''
        d.core_properties.category=''
    # Remove unused sample images/OLE objects from the supplied template.
    # Clearing visible paragraphs alone does not remove package relationships.
    used=set(d._element.xpath('//@r:embed | //@r:id | //@r:link'))
    for rid,rel in list(d.part.rels.items()):
        if rel.reltype.rsplit('/',1)[-1] in ['image','oleObject','hyperlink'] and rid not in used:
            d.part.drop_rel(rid)
    d.save(output)
    # python-docx preserves the original extended-properties part, including
    # the template author's company and stale page/word counters.
    with ZipFile(output) as z:parts={name:z.read(name) for name in z.namelist()}
    app=E.fromstring(parts['docProps/app.xml'])
    for el in app:
        if E.QName(el).localname in ['Company','Manager','TotalTime','Words','Characters','Lines','Paragraphs','CharactersWithSpaces']:
            el.text=''
        elif E.QName(el).localname=='Pages':el.text='4'
    parts['docProps/app.xml']=E.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)

def font_ppt(run,size=17,bold=False,color='222222',sup=False):
    run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    pr=run._r.get_or_add_rPr()
    for family in ['latin','ea','cs']:
        item=PX('a:'+family);item.set('typeface',FONT);pr.append(item)
    if sup:pr.set('baseline','30000')

def text_ppt(slide,text,x,y,w,h,size=17,bold=False,color='222222'):
    sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name='text:'+text[:45]
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.015);tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(3);p.line_spacing=1.05
        for bit in re.split(r'(\\\(.*?\\\))',line):
            if bit.startswith('\\('):put_powerpoint(p,bit[2:-2],size)
            elif bit:font_ppt(p.add_run(),size,bold,color);p.runs[-1].text=bit
    return sh

def math_ppt(slide,latex,x,y,w,h=.55,size=20):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    sh.name='equation:'+latex
    tf=sh.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=0;tf.margin_bottom=0
    put_powerpoint(tf.paragraphs[0],latex,size)
    return sh

def table_ppt(slide,rows,x,y,w,h,widths=None,size=14):
    sh=slide.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(w),Inches(h))
    sh.name='table'
    if widths:
        for c,width in zip(sh.table.columns,widths):c.width=Inches(width)
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            cell=sh.table.cell(i,j);cell.margin_left=Inches(.1);cell.margin_right=Inches(.08)
            cell.margin_top=Inches(.08);cell.margin_bottom=Inches(.04)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('171717' if i==0 else ('F5F5F5' if i%2==0 else 'FFFFFF'))
            cell.text_frame.clear()
            for k,line in enumerate(str(value).split('\n')):
                p=cell.text_frame.paragraphs[0] if k==0 else cell.text_frame.add_paragraph()
                p.space_after=Pt(2);p.line_spacing=1.06
                for bit in re.split(r'(\\\(.*?\\\))',line):
                    if bit.startswith('\\('):put_powerpoint(p,bit[2:-2],size,'FFFFFF' if i==0 else '222222')
                    elif bit:r=p.add_run();r.text=bit;font_ppt(r,size,i==0,'FFFFFF' if i==0 else '222222')
    return sh

def picture(slide,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    return slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))

def deck(content,output):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    for n,spec in enumerate(content['slides'],1):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        bg=s.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(255,255,255)
        text_ppt(s,f'{n:02d} | '+spec['title'],.6,.35,12.1,.63,28,True)
        refs=spec.get('references',[])
        if refs:
            sh=text_ppt(s,'',12.25,.38,.5,.35,9)
            r=sh.text_frame.paragraphs[0].add_run();r.text=' '.join(str(i+1)+')' for i in range(len(refs)));font_ppt(r,9,sup=True)
        line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.6),Inches(1.055),Inches(12.15),Inches(.012))
        line.fill.solid();line.fill.fore_color.rgb=RGBColor(30,30,30);line.line.fill.background()
        text_ppt(s,spec.get('intro',''),.65,1.28,12,.57,14,color='555555')
        for el in spec.get('elements',[]):
            kind=el['type'];x,y,w,h=el.get('box',[.65,2,12,4.5])
            if kind=='text':text_ppt(s,el['text'],x,y,w,h,el.get('size',17),el.get('bold',False))
            elif kind=='equation':math_ppt(s,el['latex'],x,y,w,h,el.get('size',20))
            elif kind=='table':table_ppt(s,el['rows'],x,y,w,h,el.get('widths'),el.get('size',14))
            elif kind=='image':picture(s,HERE/el['path'],x,y,w,h)
            elif kind=='box':
                shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('F4F4F4')
                shape.line.color.rgb=RGBColor.from_string('CCCCCC')
                text_ppt(s,el['text'],x+.15,y+.14,w-.3,h-.25,el.get('size',17),el.get('bold',False))
            else:raise ValueError(kind)
        if refs:
            y=6.78 if len(refs)>1 else 6.94
            foot=text_ppt(s,'\n'.join(f'{i+1}) {r}' for i,r in enumerate(refs)),.65,y,11.1,7.40-y,7.6,color='777777')
            for p in foot.text_frame.paragraphs:
                p.line_spacing=Pt(8.2);p.space_after=Pt(1)
        text_ppt(s,f'{n:02d} / {len(content["slides"]):02d}',12,7.1,.7,.23,8,color='777777')
        s.notes_slide.notes_text_frame.text=spec.get('notes','')
    prs.core_properties.title=content['title_ko'];prs.core_properties.author=content.get('author_line','')
    prs.save(output)

def preview_equations(source, dest):
    """LibreOffice 6 cannot render a14:m; rasterize only a disposable copy."""
    import matplotlib;matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    prs=Presentation(source)
    for s in prs.slides:
        for sh in list(s.shapes):
            if not sh._element.xpath(".//*[local-name()='oMath']"):continue
            if not sh.name.startswith('equation:'):
                # Preserve native subscripts/superscripts in the disposable
                # preview. Delivery keeps the original Office Math objects.
                for para in sh._element.xpath(".//*[local-name()='p' and namespace-uri()='http://schemas.openxmlformats.org/drawingml/2006/main']"):
                    for wrap in list(para.xpath("./*[local-name()='m']")):
                        def linear(el,baseline=0):
                            name=E.QName(el).localname
                            if name=='r' and E.QName(el).namespace.endswith('/math'):
                                value=''.join(el.xpath("./*[local-name()='t']/text()"))
                                r=PX('a:r')
                                props=el.xpath(".//*[local-name()='rPr' and namespace-uri()='http://schemas.openxmlformats.org/drawingml/2006/main']")
                                pr=deepcopy(props[0]) if props else PX('a:rPr')
                                if baseline:pr.set('baseline',str(baseline));pr.set('sz',str(int(int(pr.get('sz','1400'))*.78)))
                                r.append(pr);t=PX('a:t');t.text=value;r.append(t);wrap.addprevious(r)
                            elif name=='acc':
                                base=deepcopy(el.find('{http://schemas.openxmlformats.org/officeDocument/2006/math}e'))
                                chars=el.xpath("./*[local-name()='accPr']/*[local-name()='chr']/@*[local-name()='val']")
                                accent=chars[0] if chars else '\u0302'
                                accent={'˙':'\u0307','·':'\u0307','^':'\u0302','ˆ':'\u0302','¯':'\u0304'}.get(accent,accent)
                                texts=base.xpath(".//*[local-name()='t']")
                                if texts:texts[-1].text=(texts[-1].text or '')+accent
                                linear(base,baseline)
                            elif name in ['sSub','sSup','sSubSup']:
                                for ch in el:
                                    role=E.QName(ch).localname
                                    if role=='e':linear(ch,baseline)
                                    elif role=='sub':linear(ch,-25000)
                                    elif role=='sup':linear(ch,30000)
                            elif name not in ['rPr','oMathParaPr','sSubPr','sSupPr','sSubSupPr']:
                                for ch in el:linear(ch,baseline)
                        linear(wrap);para.remove(wrap)
                continue
            latex=sh.name.split(':',1)[1]
            sizes=sh._element.xpath(".//*[local-name()='rPr' and namespace-uri()='http://schemas.openxmlformats.org/drawingml/2006/main']/@sz")
            font_size=float(sizes[0])/100 if sizes else 20
            fig=plt.figure(figsize=(12,1.2),dpi=160)
            artist=fig.text(.01,.4,'$'+latex+'$',fontsize=font_size)
            fig.canvas.draw();bbox=artist.get_window_extent().transformed(fig.dpi_scale_trans.inverted()).expanded(1.05,1.15)
            buf=io.BytesIO();fig.savefig(buf,format='png',bbox_inches=bbox,transparent=True);plt.close(fig);buf.seek(0)
            w,h=bbox.width,bbox.height;factor=min(1.,sh.width/914400/w,sh.height/914400/h)
            pw,ph=Inches(w*factor),Inches(h*factor)
            pic=s.shapes.add_picture(buf,sh.left,sh.top+(sh.height-ph)//2,pw,ph)
            sh._element.addprevious(pic._element);sh._element.getparent().remove(sh._element)
    prs.save(dest)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--content',type=Path,default=HERE/'private/content.json')
    ap.add_argument('--templates',type=Path,default=Path('/home/jihun/Downloads'))
    ap.add_argument('--output',type=Path,default=HERE/'rendered');args=ap.parse_args()
    c=json.loads(args.content.read_text());args.output.mkdir(parents=True,exist_ok=True)
    paper(c,args.templates/'서면 심사용 논문 양식.docx',args.output/'서면심사용_블라인드_검토본_v1.docx',True)
    paper(c,args.templates/'논문양식.docx',args.output/'프로시딩_검토본_v1.docx',False)
    deck(c,args.output/'프로젝트_결과공유_검토본_v1.pptx')
    preview_equations(args.output/'프로젝트_결과공유_검토본_v1.pptx',args.output/'검수용_수식미리보기_v1.pptx')
    print('built papers and',len(c['slides']),'slides')

if __name__=='__main__':main()
